"""PowerPoint (.pptx) read/write tools."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from miqi.agent.tools.base import Tool
from miqi.agent.tools.filesystem import _persist_tracked_file
from miqi.documents.path_utils import (
    enforce_boundary,
    ensure_suffix,
    raw_output_path,
    resolve_output_path,
    resolve_read_path,
)

logger = logging.getLogger(__name__)


class PptxReadTool(Tool):
    """Read the content of a PowerPoint (.pptx) presentation."""

    name = "pptx_read"
    description = "Read and extract text content from a PowerPoint (.pptx) file."

    def __init__(
        self,
        workspace: Path | None = None,
        allowed_dir: Path | None = None,
    ):
        self._workspace = workspace
        self._allowed_dir = allowed_dir

    @property
    def parameters(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_path": {
                    "type": "string",
                    "description": "Path to the .pptx file to read. Alias for filename.",
                },
                "filename": {
                    "type": "string",
                    "description": "Filename or relative path to the .pptx file.",
                },
                "path": {
                    "type": "string",
                    "description": "Path to the .pptx file to read. Alias for filename.",
                },
            },
            "anyOf": [
                {"required": ["filename"]},
                {"required": ["file_path"]},
                {"required": ["path"]},
            ],
        }

    async def execute(self, **kwargs: Any) -> str:
        _sess_key = kwargs.pop("_session_key", None)
        raw_path = raw_output_path(kwargs)
        if not raw_path.strip():
            return "Error: 必须提供 filename"
        try:
            file_path = resolve_output_path(
                raw_path, self._workspace, self._allowed_dir,
            )
            file_path = ensure_suffix(file_path, ".pptx")
            enforce_boundary(file_path, self._allowed_dir, self._workspace)
        except PermissionError as e:
            return f"Error: 权限被拒绝：{e}"
        except ValueError as e:
            return f"Error: {e}"
        if not file_path.exists():
            return f"Error: 文件不存在：{file_path}"
        try:
            from pptx import Presentation
            prs = Presentation(str(file_path))
            lines = [f"Presentation: {len(prs.slides)} slides", ""]
            for i, slide in enumerate(prs.slides, 1):
                lines.append(f"## Slide {i}")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text:
                                lines.append(text)
                lines.append("")
            return "\n".join(lines)
        except Exception as e:
            return f"Error reading {file_path.name}: {e}"


class CreatePptxTool(Tool):
    """Create a PowerPoint (.pptx) presentation."""

    name = "create_pptx"
    description = (
        "Create a PowerPoint (.pptx) presentation in the workspace files directory. "
        "Supports multiple slides with titles, bullets, body text, and images. "
        "Each slide's image_path must resolve inside the session files directory "
        "(or a user-authorized directory); an image outside those roots is "
        "skipped, not embedded, and reported in the result."
    )

    def __init__(
        self,
        workspace: Path | None = None,
        allowed_dir: Path | None = None,
        allow_user_roots: bool = False,
    ):
        self._workspace = workspace
        self._allowed_dir = allowed_dir
        self._allow_user_roots = allow_user_roots

    @property
    def parameters(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_path": {
                    "type": "string",
                    "description": "Path for the output .pptx file. Alias for filename.",
                },
                "filename": {
                    "type": "string",
                    "description": "Filename or relative path for the output .pptx file",
                },
                "path": {
                    "type": "string",
                    "description": "Path for the output .pptx file. Alias for filename.",
                },
                "slides": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string"},
                            "subtitle": {"type": "string"},
                            "content": {
                                "description": "Slide body text or an array of body lines.",
                            },
                            "bullets": {"type": "array", "items": {"type": "string"}},
                            "image_path": {"type": "string"},
                        },
                    },
                    "description": "List of slides with title, content, bullets, or image_path",
                },
            },
            "anyOf": [
                {"required": ["filename"]},
                {"required": ["file_path"]},
                {"required": ["path"]},
            ],
        }

    async def execute(self, **kwargs: Any) -> str:
        from pptx import Presentation
        from pptx.util import Inches

        _sess_key = kwargs.pop("_session_key", None)
        user_roots = kwargs.pop("_user_roots", None)
        raw_path = raw_output_path(kwargs)
        slides = kwargs.get("slides") or []
        if not raw_path.strip():
            return "Error: 必须提供 filename"

        try:
            file_path = resolve_output_path(
                raw_path, self._workspace, self._allowed_dir,
            )
            file_path = ensure_suffix(file_path, ".pptx")
            enforce_boundary(file_path, self._allowed_dir, self._workspace)
        except PermissionError as e:
            return f"Error: 权限被拒绝：{e}"
        except ValueError as e:
            return f"Error: {e}"
        if not slides:
            return "Error: 必须提供 slides"

        skipped_images: list[str] = []
        try:
            prs = Presentation()
            for slide_data in slides:
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                title = slide.shapes.title
                if title:
                    title.text = str(slide_data.get("title", ""))
                body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
                if body and hasattr(body, "text_frame"):
                    text_frame = body.text_frame
                    text_frame.clear()
                    content_items: list[Any] = []
                    if slide_data.get("subtitle"):
                        content_items.append(slide_data["subtitle"])
                    content = slide_data.get("content")
                    if isinstance(content, list):
                        content_items.extend(content)
                    elif content:
                        content_items.append(content)
                    bullets = slide_data.get("bullets") or []
                    for item_index, item in enumerate(content_items):
                        if item_index == 0:
                            text_frame.text = str(item)
                        else:
                            paragraph = text_frame.add_paragraph()
                            paragraph.text = str(item)
                            paragraph.level = 0
                    for bullet in bullets:
                        paragraph = text_frame.add_paragraph()
                        paragraph.text = str(bullet)
                        paragraph.level = 0
                image_path = slide_data.get("image_path")
                if image_path:
                    try:
                        resolved = resolve_read_path(
                            str(image_path),
                            self._workspace,
                            self._allowed_dir,
                            user_roots,
                            self._allow_user_roots,
                        )
                        slide.shapes.add_picture(
                            str(resolved),
                            Inches(float(slide_data.get("image_left", 5.5))),
                            Inches(float(slide_data.get("image_top", 1.5))),
                            width=Inches(float(slide_data.get("image_width", 3.0))),
                        )
                    except Exception as e:
                        # Out-of-bounds, missing, or not-an-image: skip this
                        # slide's picture and keep building the deck
                        # (issue #1005 节 2).
                        logger.warning(
                            "create_pptx: 跳过图片 %s：%s", image_path, e,
                        )
                        skipped_images.append(f"{image_path}（{e}）")

            file_path.parent.mkdir(parents=True, exist_ok=True)
            prs.save(str(file_path))
            _persist_tracked_file(self._workspace, file_path, op="write", session_key=_sess_key)
            result = f"Created: {file_path} ({len(slides)} slides)"
            if skipped_images:
                result += (
                    f"（跳过 {len(skipped_images)} 个图片："
                    f"{'；'.join(skipped_images)}）"
                )
            return result
        except Exception as e:
            return f"Error writing {raw_path}: {e}"


class PptxWriteTool(CreatePptxTool):
    """Backward-compatible alias for create_pptx."""

    name = "pptx_write"
    description = (
        "Create a new PowerPoint (.pptx) file. Each slide's image_path must "
        "resolve inside the session files directory (or a user-authorized "
        "directory); out-of-bounds images are skipped and reported. "
        "Prefer create_pptx for new calls."
    )
